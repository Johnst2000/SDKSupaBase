"""Genera PRÁCTICA EN CLASE. LISTVIEW GLIDE SUPABASE SDK.docx con código del proyecto Android."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parent
OUTPUT_PATH = PROJECT_ROOT / "PRÁCTICA EN CLASE. LISTVIEW GLIDE SUPABASE SDK.docx"

SRC = PROJECT_ROOT / "app" / "src" / "main"
JAVA = SRC / "java" / "com" / "uteq" / "software" / "app"


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def extract_block(source: str, start_marker: str, end_marker: str | None = None) -> str:
    start = source.find(start_marker)
    if start == -1:
        return source
    if end_marker is None:
        depth = 0
        for i in range(start, len(source)):
            ch = source[i]
            if ch == "{":
                depth += 1
            elif ch == "}":
                depth -= 1
                if depth == 0:
                    return source[start : i + 1]
        return source[start:]
    end = source.find(end_marker, start + len(start_marker))
    if end == -1:
        return source[start:]
    return source[start : end + len(end_marker)]


def main_activity2_supabase_sections() -> str:
    full = read_text(JAVA / "Utils" / "MainActivity2.kt")
    imports = extract_block(full, "import com.uteq.software.app.Services.SupabaseErrorHandler", "class MainActivity2")
    class_header = "class MainActivity2 : AppCompatActivity() {\n\n    private lateinit var spinnerSemestre: Spinner\n    private lateinit var spinnerMaterias: Spinner\n    private lateinit var listView: ListView\n\n    private val categorias by lazy { resources.getStringArray(R.array.niveles).toList() }\n    private var materiasFiltradas = listOf<Materia>()\n    private var alumnos = ArrayList<Alumno>()"
    on_create_launch = extract_block(full, "        lifecycleScope.launch {\n            cargarDatosIniciales()\n        }", "    private fun configurarSpinnerSemestre")
    cargar_datos = extract_block(full, "    private suspend fun cargarDatosIniciales()", "    private suspend fun cargarMaterias")
    cargar_materias = extract_block(full, "    private suspend fun cargarMaterias(nivel: Int)", "    private fun actualizarListaAlumnos")
    actualizar = extract_block(full, "    private fun actualizarListaAlumnos()", "}\n")
    return (
        "package com.uteq.software.app.Utils\n\n"
        "// Imports relevantes para consulta Supabase\n"
        + imports.strip()
        + "\n\n"
        + class_header
        + "\n\n    // ... onCreate y listeners ...\n\n"
        + on_create_launch.strip()
        + "\n\n"
        + cargar_datos.strip()
        + "\n\n"
        + cargar_materias.strip()
        + "\n\n"
        + actualizar.strip()
        + "\n}"
    )


def main_activity_supabase_section() -> str:
    full = read_text(JAVA / "Utils" / "MainActivity.kt")
    return extract_block(
        full,
        "        lifecycleScope.launch {",
        "        }\n    }\n}",
    )


def build_config_section() -> str:
    gradle = read_text(PROJECT_ROOT / "app" / "build.gradle.kts")
    default_config = extract_block(gradle, "    defaultConfig {", "    }\n\n    buildFeatures")
    build_features = extract_block(gradle, "    buildFeatures {", "    }\n\n    buildTypes")
    dependencies = extract_block(gradle, "dependencies {")
    return (
        "// app/build.gradle.kts — sección defaultConfig (credenciales vía local.properties)\n"
        + default_config
        + "\n\n"
        + build_features
        + "\n\n"
        + dependencies
    )


def add_title_page(doc: Document) -> None:
    for _ in range(6):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("PRÁCTICA EN CLASE")
    run.bold = True
    run.font.size = Pt(26)
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor(0x1A, 0x47, 0x7A)

    doc.add_paragraph()
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("ListView, Glide, Supabase SDK")
    r.bold = True
    r.font.size = Pt(18)
    r.font.name = "Calibri"

    sub2 = doc.add_paragraph()
    sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub2.add_run("Universidad Técnica Estatal de Quevedo (UTEQ)")
    r2.font.size = Pt(14)
    r2.font.name = "Calibri"

    sub3 = doc.add_paragraph()
    sub3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = sub3.add_run("Aplicaciones Móviles — 6° Semestre")
    r3.font.size = Pt(12)
    r3.font.name = "Calibri"

    sub4 = doc.add_paragraph()
    sub4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r4 = sub4.add_run(f"Fecha: {date.today().strftime('%d/%m/%Y')}")
    r4.font.size = Pt(11)
    r4.font.name = "Calibri"
    r4.italic = True

    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(0x1A, 0x47, 0x7A)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.bold = bold


def add_code_block(doc: Document, code: str, filename: str | None = None) -> None:
    if filename:
        cap = doc.add_paragraph()
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.space_after = Pt(4)
        r = cap.add_run(f"Archivo: {filename}")
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    for line in code.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)

    doc.add_paragraph()


def add_checklist_item(doc: Document, text: str, done: bool = True) -> None:
    mark = "[✓]" if done else "[ ]"
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(f"{mark} {text}")
    run.font.name = "Calibri"
    run.font.size = Pt(11)


def add_screenshot_placeholder(doc: Document, title: str, description: str) -> None:
    add_heading(doc, title, level=3)
    add_para(doc, description)
    box = doc.add_paragraph()
    box.paragraph_format.left_indent = Cm(1)
    box.paragraph_format.space_before = Pt(12)
    box.paragraph_format.space_after = Pt(24)
    run = box.add_run("[ Insertar captura de pantalla aquí ]")
    run.italic = True
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)


def build_document() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    add_title_page(doc)

    # 1. Descripción y cumplimiento de rúbrica
    add_heading(doc, "1. Descripción y cumplimiento de rúbrica")
    add_para(
        doc,
        "Aplicación Android que integra el SDK de Supabase (PostgREST) con ListView personalizado, "
        "carga de imágenes con Glide y consultas asíncronas mediante Kotlin Coroutines. "
        "MainActivity2 es la pantalla principal (launcher) con spinners de Semestre/Materias y lista de alumnos. "
        "MainActivity muestra una vista alternativa con el listado en texto plano.",
    )
    add_para(doc, "Checklist de requisitos de la rúbrica:", bold=True)
    add_checklist_item(
        doc,
        "Clase de datos Alumno (@Serializable) con campos id, nombres, correo, telefono, foto",
    )
    add_checklist_item(doc, "AlumnoAdapter personalizado para ListView con Glide (foto circular)")
    add_checklist_item(
        doc,
        "Consulta asíncrona a Supabase con orden alfabético (order nombres ASC) en MainActivity2 y MainActivity",
    )
    add_checklist_item(
        doc,
        "SupabaseManager singleton con credenciales desde BuildConfig (local.properties, sin claves en código)",
    )
    add_checklist_item(doc, "Layouts XML: activity_main2.xml (principal) e item_alumno.xml (ítem de lista)")
    add_checklist_item(doc, "strings.xml con array de niveles/semestres (Primero … Séptimo)")
    add_checklist_item(
        doc,
        "build.gradle.kts con dependencias Supabase, Glide, Serialization y buildConfig habilitado",
    )
    add_checklist_item(doc, "README con stack tecnológico, estructura, requisitos e instrucciones de ejecución")

    # 2. Código clase Alumno
    add_heading(doc, "2. Código clase Alumno")
    add_code_block(doc, read_text(JAVA / "Models" / "Alumno.kt"), "Models/Alumno.kt")

    # 3. Código AlumnoAdapter
    add_heading(doc, "3. Código AlumnoAdapter")
    add_code_block(doc, read_text(JAVA / "Adapters" / "AlumnoAdapter.kt"), "Adapters/AlumnoAdapter.kt")

    # 4. Código consulta Supabase async
    add_heading(doc, "4. Código consulta Supabase (asíncrono)")
    add_para(
        doc,
        "Secciones relevantes de MainActivity2.kt: carga inicial de alumnos ordenados alfabéticamente, "
        "consulta de materias filtradas por nivel y actualización del ListView.",
    )
    add_code_block(doc, main_activity2_supabase_sections(), "Utils/MainActivity2.kt (secciones relevantes)")

    add_para(doc, "Consulta equivalente en MainActivity.kt (vista de texto):", bold=True)
    add_code_block(doc, main_activity_supabase_section(), "Utils/MainActivity.kt (lifecycleScope.launch)")

    add_para(doc, "Cliente Supabase (singleton):", bold=True)
    add_code_block(doc, read_text(JAVA / "Services" / "SupaBaseManager.kt"), "Services/SupaBaseManager.kt")

    # 5. Configuración credenciales
    add_heading(doc, "5. Configuración de credenciales")
    add_para(
        doc,
        "Las credenciales NO se incluyen en el repositorio. Se definen en local.properties "
        "(ignorado por .gitignore) y Gradle las expone como campos BuildConfig.",
    )
    add_para(doc, "Ejemplo de local.properties:", bold=True)
    add_code_block(doc, read_text(PROJECT_ROOT / "local.properties.example"), "local.properties.example")

    add_para(doc, "Sección buildConfig y dependencias en app/build.gradle.kts:", bold=True)
    add_code_block(doc, build_config_section(), "app/build.gradle.kts")

    # 6. Layouts XML principales
    add_heading(doc, "6. Layouts XML principales")
    add_heading(doc, "6.1 activity_main2.xml", level=2)
    add_code_block(
        doc,
        read_text(SRC / "res" / "layout" / "activity_main2.xml"),
        "res/layout/activity_main2.xml",
    )
    add_heading(doc, "6.2 item_alumno.xml", level=2)
    add_code_block(doc, read_text(SRC / "res" / "layout" / "item_alumno.xml"), "res/layout/item_alumno.xml")

    # 7. Strings/values
    add_heading(doc, "7. Strings y values")
    add_para(doc, "Archivo strings.xml (sin colors.xml en esta sección):")
    add_code_block(doc, read_text(SRC / "res" / "values" / "strings.xml"), "res/values/strings.xml")

    # 8. Capturas
    add_heading(doc, "8. Capturas de pantalla")
    add_para(doc, "Espacios reservados para pegar imágenes del emulador o dispositivo físico.")
    add_screenshot_placeholder(
        doc,
        "8.1 Pantalla principal MainActivity2",
        "Logo UTEQ, spinners de Semestre y Materias, ListView con alumnos.",
    )
    add_screenshot_placeholder(
        doc,
        "8.2 Ítem de alumno con Glide",
        "Foto circular, nombre en mayúsculas, correo y teléfono.",
    )
    add_screenshot_placeholder(
        doc,
        "8.3 Filtro por semestre y materias",
        "Dropdowns cargados desde Supabase filtrando por nivel.",
    )
    add_screenshot_placeholder(
        doc,
        "8.4 MainActivity (lista en texto)",
        "Vista secundaria con alumnos en EditText de solo lectura.",
    )
    add_screenshot_placeholder(
        doc,
        "8.5 Diálogo de error Supabase",
        "Captura opcional del MaterialAlertDialogBuilder ante RestException.",
    )

    # 9. URL repositorio GitHub
    add_heading(doc, "9. URL del repositorio GitHub")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("[ Insertar URL del repositorio GitHub aquí ]")
    run.italic = True
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    # 10. Instrucciones de instalación
    add_heading(doc, "10. Instrucciones de instalación")
    steps = [
        "Clonar o descargar el repositorio SDKSupaBase en el equipo local.",
        "Copiar local.properties.example como local.properties en la raíz del proyecto.",
        "Editar local.properties: configurar sdk.dir y las variables SUPABASE_URL y SUPABASE_KEY (placeholders, no subir claves reales).",
        "Abrir la carpeta SDKSupaBase en Android Studio (File → Open) y esperar Gradle Sync.",
        "Verificar JDK 11, minSdk 26 y conexión a Internet (permiso INTERNET en AndroidManifest).",
        "Crear o seleccionar un emulador (API 26+) o dispositivo físico con depuración USB.",
        "Ejecutar Run ▶ (Shift + F10). La app abre MainActivity2 como pantalla principal.",
    ]
    for i, step in enumerate(steps, 1):
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(step)
        run.font.name = "Calibri"
        run.font.size = Pt(11)

    return doc


def main() -> None:
    doc = build_document()
    doc.save(str(OUTPUT_PATH))
    print(f"Documento generado: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
